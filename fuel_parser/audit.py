"""Audit workbook: reconciliation, mapping status, review list and row log (design F.5, G, N.2)."""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import file_sha256
from .model import RecordStatus, Severity
from .writer import set_text

if TYPE_CHECKING:
    from .pipeline import RunOptions, RunResult

_BOLD = Font(bold=True)
_REVIEW_SEVERITIES = (Severity.ERROR, Severity.HELD, Severity.WARNING)


def _sheet(wb, title: str, header: list[str]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = _BOLD
    ws.freeze_panes = "A2"
    _next_row[id(ws)] = 2
    return ws


_next_row: dict[int, int] = {}


def _append(ws, values: list[object]) -> None:
    # Track the next row ourselves: openpyxl's max_row scans every cell (quadratic on big sheets).
    r = _next_row[id(ws)] if id(ws) in _next_row else ws.max_row + 1
    _next_row[id(ws)] = r + 1
    for c, v in enumerate(values, start=1):
        cell = ws.cell(r, c)
        if isinstance(v, (datetime, int, float)) and not isinstance(v, bool):
            cell.value = v
        elif v is not None:
            set_text(cell, str(v))


def mapping_status_rows(result: "RunResult") -> list[list[object]]:
    cfg = result.cfg
    per_leader: dict[str, Counter] = defaultdict(Counter)
    for rec in result.records:
        per_leader[rec.sales_leader][rec.status.value] += 1
    rows = []
    for leader in sorted(per_leader):
        team = cfg.teams.lookup(leader) if leader else None
        c = per_leader[leader]
        held = sum(v for k, v in c.items() if k != RecordStatus.EXPORTED.value)
        rows.append([
            leader or "(blank)",
            team.status if team else "UNRESOLVED (not in teams.yaml)",
            team.slot if team and team.exportable else None,
            team.describe() if team else "No entry",
            c.get(RecordStatus.EXPORTED.value, 0),
            held,
        ])
    return rows


def write_audit(path: Path, result: "RunResult", opts: "RunOptions") -> None:
    cfg = result.cfg
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    _next_row[id(ws)] = 1
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 70
    by_status = Counter(r.status.value for r in result.records)
    owner_exported = sum(1 for r in result.exported if r.slot_status == "OWNER_APPROVED_UNVERIFIED")
    lines: list[list[object]] = [
        ["FUEL Response Parser — audit report", ""],
        ["Mode", "DRY RUN (no import file, ledger unchanged)" if result.dry_run else "EXPORT"],
        ["Batch id", result.batch_id],
        ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Source file", Path(opts.source).name],
        ["Source sha256", result.source_sha256],
        ["Response sheet used", result.sheet.sheet_name],
        ["Other sheets ignored", ", ".join(result.sheet.ignored_sheets) or "(none)"],
        ["Template file", Path(opts.template).name],
        ["Template sha256", file_sha256(opts.template)],
        ["Mapping profile sha256", cfg.profile.sha256],
        ["Header signature", result.cmap.signature],
        ["Vocabulary sha256", cfg.vocab.sha256],
        ["Validation policy sha256", cfg.policy.sha256],
        ["Teams config sha256", cfg.teams.sha256],
        ["Seller directory", f"{cfg.sellers.source} ({len(cfg.sellers.emails)} approved) {cfg.sellers.sha256}"],
        ["Team filter", ", ".join(opts.team_filter) if opts.team_filter else "(all teams)"],
        ["", ""],
    ]
    if result.previous_batches:
        lines.append(["WARNING: this source file was exported before",
                      "; ".join(f"{b.batch_id} ({b.records_exported} exported)" for b in result.previous_batches)])
        lines.append(["", ""])
    lines += [
        ["Headline",
         f"Exported: {len(result.exported) - owner_exported} verified · {owner_exported} owner-approved (unverified) · "
         f"Held for review: {by_status.get('REVIEW', 0)} · Duplicates held: {by_status.get('DUPLICATE_HELD', 0)} · "
         f"Already imported: {by_status.get('ALREADY_IMPORTED', 0)} · Excluded by filter: {by_status.get('EXCLUDED_BY_FILTER', 0)}"],
        ["", ""],
        ["Reconciliation", "Value"],
    ]
    lines += [[k, v] for k, v in result.counts.items()]
    lines += [["", ""], ["Reconciliation check", "Result"]]
    lines += [[name, "PASS" if ok else "FAIL"] for name, ok, _ in result.checks]
    issue_counts = Counter((i.code, i.severity.value) for r in result.records for i in r.issues)
    issue_counts.update((i.code, i.severity.value) for i in result.response_issues)
    lines += [["", ""], ["Issue code (severity)", "Count"]]
    lines += [[f"{code} ({sev})", n] for (code, sev), n in sorted(issue_counts.items())]
    for line in lines:
        _append(ws, line)
    for row in ws.iter_rows():
        if row[0].value in ("Reconciliation", "Reconciliation check", "Issue code (severity)", "Headline"):
            row[0].font = row[1].font = _BOLD

    ms = _sheet(wb, "MappingStatus", ["SALES LEADER", "Status", "Slot", "Evidence / approval", "Records exported", "Records held"])
    for row in mapping_status_rows(result):
        _append(ms, row)

    dest_fields = [f for f in cfg.profile.dest_columns if not f.startswith("seller_") and f != "id"]
    rv = _sheet(wb, "Review", ["Status", "Reasons", "Action", "Source row", "Block", "SALES LEADER", "Slot status",
                               "Seller name", "Seller e-mail", "Lead source (received)", "Duplicate group"] + dest_fields)
    for rec in sorted(result.records, key=lambda r: (r.source_row, r.block_index or 0)):
        flagged = [i for i in rec.issues if i.severity in _REVIEW_SEVERITIES]
        if rec.status == RecordStatus.EXPORTED and not flagged:
            continue
        if rec.status in (RecordStatus.ALREADY_IMPORTED, RecordStatus.EXCLUDED_BY_FILTER) and not flagged:
            continue
        codes = list(dict.fromkeys(f"{i.code}" + (f"[{i.field}]" if i.field else "") for i in flagged))
        actions = list(dict.fromkeys(cfg.policy.action.get(i.code, "") for i in flagged if cfg.policy.action.get(i.code)))
        _append(rv, [rec.status.value, "; ".join(codes), " ".join(actions), rec.source_row, rec.block_index,
                     rec.sales_leader, rec.slot_status, rec.meta.get("seller_name"), rec.meta.get("seller_email"),
                     rec.lead_source_original, rec.duplicate_group] + [rec.values.get(f) for f in dest_fields])
    for i in result.response_issues:
        _append(rv, ["REVIEW", i.code, cfg.policy.action.get(i.code, ""), i.source_row, None, None, None, None, None, None, None])

    iss = _sheet(wb, "Issues", ["Source row", "Block", "Record key", "Code", "Severity", "Field", "Message"])
    for rec in result.records:
        for i in rec.issues:
            _append(iss, [rec.source_row, rec.block_index, rec.src_key[:12], i.code, i.severity.value, i.field, i.message])
    for i in result.response_issues:
        _append(iss, [i.source_row, None, "", i.code, i.severity.value, i.field, i.message])

    rl = _sheet(wb, "RowLog", ["Source row", "Block", "Status", "Output Id", "Record key", "Lead source (received)", "Issue codes"])
    entries = [(r.source_row, r.block_index or 0, [r.source_row, r.block_index, r.status.value, r.dest_id, r.src_key,
                r.lead_source_original, ", ".join(dict.fromkeys(r.codes()))]) for r in result.records]
    entries += [(i.source_row, 0, [i.source_row, None, "REVIEW", None, "", "", i.code]) for i in result.response_issues]
    for _, _, row in sorted(entries, key=lambda e: (e[0], e[1])):
        _append(rl, row)

    mp = _sheet(wb, "Mapping", ["Column", "Source header", "Result", "Target", "Note"])
    for p in result.cmap.preview:
        _append(mp, [p.letter, p.header, p.kind, p.target, p.note])

    um = _sheet(wb, "Unmapped", ["Column", "Source header", "Note"])
    for p in result.cmap.preview:
        if p.kind in ("IGNORED_EMPTY",) or p.target == "ignore":
            _append(um, [p.letter, p.header, p.note])

    if opts.team_sheet_check and result.sheet.path.suffix.lower() in (".xlsx", ".xlsm"):
        team_sheet_check(wb, result)  # opt-in reference comparison; never an input (H-33)
    wb.save(path)


# -- optional comparison with the derived team sheets (M.2) -------------------------------

_TEAM_FIELDS = {  # Leads Log header -> how to read our record
    "SELLER EMAIL ADDRESS": lambda r: r.values.get("email"),
    "AAD": lambda r: r.sales_leader,
    "AS": lambda r: r.meta.get("seller_name"),
    "SURNAME (CLIENT)": lambda r: r.values.get("surname"),
    "FIRST NAME (CLIENT)": lambda r: r.values.get("first_name"),
    "MIDDLE NAME (CLIENT)": lambda r: r.values.get("middle_name"),
    "MOBILE NO.": lambda r: r.values.get("mobile"),
    "EMAIL ADDRESS": lambda r: r.values.get("client_email"),
    "RESIDENCE ADDRESS": lambda r: r.values.get("residence"),
    "PROJECT (FIRST CHOICE)": lambda r: r.values.get("project_first"),
    "PROJECT (SECOND CHOICE)": lambda r: r.values.get("project_second"),
    "REMARKS": lambda r: r.meta.get("remarks_original"),
    "LEAD SOURCE": lambda r: r.lead_source_original,
    "SOURCE OF AWARENESS": lambda r: r.values.get("awareness"),
}
_MULTI = {"PROJECT (FIRST CHOICE)", "PROJECT (SECOND CHOICE)", "SOURCE OF AWARENESS"}


def _norm(v: object, header: str, null_tokens: set[str]) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if s.upper() in {t.upper() for t in null_tokens}:
        return ""  # team sheets upper-case values, so placeholders are compared case-insensitively
    s = s.upper()
    if header in _MULTI:
        s = re.sub(r"[;,\s]", "", s)
    return s


def _ts_key(v: object) -> int | None:
    return round(v.timestamp()) if isinstance(v, datetime) else None


def team_sheet_check(wb, result: "RunResult") -> None:
    """Compare our records with the team sheets' cached values (reference only, case-insensitive).

    SPECIFIC / DISTINCT are excluded (I.5 X7). Values are not copied into the report.
    """
    src = openpyxl.load_workbook(result.sheet.path, read_only=True, data_only=True)
    try:
        sheets = []
        for name in src.sheetnames:
            if name == result.sheet.sheet_name:
                continue
            rows = list(src[name].iter_rows(max_row=10, values_only=True))
            for idx, row in enumerate(rows):
                labels = [str(c).strip().upper() if c is not None else "" for c in row]
                if "TIMESTAMP" in labels and "AAD" in labels and "AS" in labels:
                    sheets.append((name, idx + 1, labels))
                    break
        if not sheets:
            return
        ours: dict[tuple, list] = defaultdict(list)
        for rec in result.records:
            k = (_ts_key(rec.timestamp), _norm(rec.values.get("surname"), "", set()))
            ours[k].append(rec)
        nulls = result.cfg.vocab.null_tokens
        stats: dict[str, Counter] = defaultdict(Counter)
        mismatches, unmatched = [], 0
        compared = 0
        for name, hrow, labels in sheets:
            col = {lab: i for i, lab in enumerate(labels) if lab}
            for n, row in enumerate(src[name].iter_rows(min_row=hrow + 1, values_only=True), start=hrow + 1):
                row = list(row)
                ts = row[col["TIMESTAMP"]] if col.get("TIMESTAMP") is not None and col["TIMESTAMP"] < len(row) else None
                if ts is None:
                    continue
                sur_i = col.get("SURNAME (CLIENT)")
                # Team sheets hold timestamps as float day serials, which drift by up to ±1 s.
                base, sur = _ts_key(ts), _norm(row[sur_i] if sur_i is not None else None, "", set())
                recs = None
                if base is not None:
                    recs = next((ours[(base + d, sur)] for d in (0, -1, 1) if (base + d, sur) in ours), None)
                if not recs:
                    unmatched += 1
                    continue
                rec = recs[0]
                compared += 1
                for header, getter in _TEAM_FIELDS.items():
                    i = col.get(header)
                    if i is None or i >= len(row):
                        continue
                    if _norm(row[i], header, nulls) == _norm(getter(rec), header, nulls):
                        stats[header]["match"] += 1
                    else:
                        stats[header]["mismatch"] += 1
                        mismatches.append([name, n, rec.source_row, header])
        ws = _sheet(wb, "TeamSheetCheck", ["Field (reference: team sheets, case-insensitive)", "Matched", "Mismatched"])
        _append(ws, ["Team-sheet rows compared", compared, None])
        _append(ws, ["Team-sheet rows with no matching record", unmatched, None])
        for header in _TEAM_FIELDS:
            _append(ws, [header, stats[header]["match"], stats[header]["mismatch"]])
        ws2 = _sheet(wb, "TeamSheetMismatches", ["Team sheet", "Team-sheet row", "Source row", "Field"])
        for m in mismatches:
            _append(ws2, m)
    finally:
        src.close()
