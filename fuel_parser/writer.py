"""Write the import-ready MASTERLIST and the optional Leads Log (design D.6, F.5, I.2)."""

from __future__ import annotations

import functools
import hashlib
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .config import Profile
from .model import Record


class TemplateError(ValueError):
    """Raised when the destination template does not match the approved schema."""


def _header_digest(header: list[object]) -> str:
    return hashlib.sha256("\x1f".join("" if v is None else str(v) for v in header).encode("utf-8")).hexdigest()


@functools.lru_cache(maxsize=4)
def _template_header_cached(path: str, size: int, mtime: float, preferred: str, digest: str):
    wb = openpyxl.load_workbook(path, read_only=False)
    try:
        # The approved sheet name first, then any other sheet: the template is identified by its exact header.
        names = [preferred] * (preferred in wb.sheetnames) + [n for n in wb.sheetnames if n != preferred]
        for name in names:
            ws = wb[name]
            header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            while header and header[-1] is None:
                header.pop()
            if _header_digest(header) == digest:
                widths = {get_column_letter(i): ws.column_dimensions[get_column_letter(i)].width for i in range(1, len(header) + 1)}
                return tuple(header), tuple(widths.items())
    finally:
        wb.close()
    raise TemplateError(
        "No sheet in the template file has the approved 35-column MASTERLIST header. "
        "No output was written. Check that the correct template was supplied.")


def read_template_header(template: Path, profile: Profile) -> tuple[list[object], dict[str, float]]:
    """Return the template's exact header cells and column widths, after the schema check (T43).

    The sheet is found by its byte-exact header, so a template whose sheet is not named MASTERLIST is accepted.
    """
    st = Path(template).stat()
    header, widths = _template_header_cached(str(Path(template).resolve()), st.st_size, st.st_mtime,
                                             profile.dest["sheet"], profile.dest["template_header_sha256"])
    return list(header), dict(widths)


def set_text(cell, value: object) -> bool:
    """Write a value as a literal string. Returns True when formula-like text was escaped."""
    cell.value = value
    cell.number_format = "@"
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # openpyxl would otherwise store it as a formula (D.6)
        return True
    return False


def write_masterlist(path: Path, header: list[object], widths: dict[str, float], records: list[Record],
                     profile: Profile) -> list[Record]:
    """Write exported records. Returns records whose text needed formula escaping."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = profile.dest["sheet"]
    for c, h in enumerate(header, start=1):
        ws.cell(1, c).value = h  # byte-exact header (AC8)
    for letter, w in widths.items():
        if w:
            ws.column_dimensions[letter].width = w
    dt_cols = set(profile.dest.get("datetime_columns") or [])
    dt_fmt = profile.dest.get("datetime_format", "m/d/yy h:mm")
    escaped = []
    for r, rec in enumerate(records, start=2):
        for c, fid in enumerate(profile.dest_columns, start=1):
            val = rec.values.get(fid)
            cell = ws.cell(r, c)
            if fid == "id":
                cell.value = rec.dest_id
            elif fid in dt_cols:
                cell.value = val
                cell.number_format = dt_fmt
            elif val is None:
                cell.number_format = "@"
            elif set_text(cell, str(val)):
                escaped.append(rec)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(1, len(records) + 1)}"
    wb.save(path)
    return escaped


def leads_log_row(values: dict, slot: int | None, profile: Profile) -> list[object]:
    """Project one MASTERLIST record onto the 17 Leads Log columns (I.2)."""
    ls = values.get("lead_source")
    detail = profile.lead_source_detail.get(ls) if ls else None
    specific = values.get(detail) if detail and detail != "question" else None
    return [
        values.get("completion_time"),
        values.get(f"seller_email_{slot}") if slot else None,
        values.get("sales_leader"),
        values.get(f"seller_name_{slot}") if slot else None,
        values.get("surname"), values.get("first_name"), values.get("middle_name"),
        None if values.get("mobile") is None else str(values.get("mobile")),
        values.get("client_email"), values.get("residence"),
        values.get("project_first"), values.get("project_second"), values.get("remarks"),
        ls, specific,
        None,  # LEAD SOURCE (DISTINCT): blank until a rule is approved (H-23)
        values.get("awareness"),
    ]


def slot_mismatch(values: dict, slot: int) -> bool:
    """True when the configured slot is empty while another slot holds seller data (N.2)."""
    def filled(n: int) -> bool:
        return bool(values.get(f"seller_name_{n}") or values.get(f"seller_email_{n}"))
    return not filled(slot) and any(filled(n) for n in (1, 2, 3) if n != slot)


def write_leads_log(path: Path, rows: list[list[object]], profile: Profile) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LEADS LOG"
    headers = profile.leads_log["headers"]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c).value = h
    fmt = profile.leads_log.get("timestamp_format", "mmm/dd/yyyy hh:mm:ss")
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(r, c)
            if c == 1 and isinstance(val, datetime):
                cell.value = val
                cell.number_format = fmt  # H-22: real datetime
            elif val is None:
                cell.number_format = "@"
            else:
                set_text(cell, str(val))
    wb.save(path)
