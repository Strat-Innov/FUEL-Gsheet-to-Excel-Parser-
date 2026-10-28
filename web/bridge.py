"""Browser bridge (decision H-32): runs fuel_parser inside the user's browser via Pyodide.

Files are read from the in-browser file system; nothing is uploaded. The import ledger lives in
browser storage (IndexedDB) under /ledger. Team assignments made on screen are the user's explicit
owner-approved assumptions (OWNER_APPROVED_UNVERIFIED), recorded with approver and date (N.1).
"""

from __future__ import annotations

import json
import traceback
from collections import Counter, defaultdict
from pathlib import Path

import yaml

from fuel_parser.audit import mapping_status_rows
from fuel_parser.config import ConfigError, load_teams
from fuel_parser.history import Ledger
from fuel_parser.mapping import MappingError
from fuel_parser.model import RecordStatus, Severity
from fuel_parser.pipeline import NeedsConfirmation, ReconciliationError, RunOptions, run as run_pipeline
from fuel_parser.reader import ReadError
from fuel_parser.writer import TemplateError

APP = Path("/app")
WORK = Path("/work")
LEDGER = Path("/ledger/import_history.sqlite")
_ORDER = ["ERROR", "HELD", "WARNING", "INFO"]
_KNOWN = (ConfigError, MappingError, ReadError, TemplateError, ReconciliationError, ValueError, KeyError)


def _teams_with_overrides(overrides: dict) -> Path:
    """Write a teams.yaml = shipped config + on-screen owner approvals. Returns its path."""
    teams = load_teams(APP / "config" / "teams.yaml")
    entries = []
    for e in teams.entries:
        entries.append({k: v for k, v in vars(e).items() if v not in (None, "") or k == "slot"})
    for key, ov in (overrides or {}).items():
        slot = ov.get("slot")
        if slot not in (1, 2, 3):
            continue
        approval = {"slot": slot, "status": "OWNER_APPROVED_UNVERIFIED", "evidence_type": "business_rule",
                    "approved_by": str(ov.get("approved_by") or "").strip(), "approved_on": str(ov.get("approved_on") or "")}
        match = teams.lookup(key)
        if match is not None and match.exportable:
            continue  # a VERIFIED or approved entry in the shipped config wins
        if match is not None:
            for ent in entries:
                if ent["label_prefix"] == match.label_prefix:
                    ent.update(approval)
                    ent.pop("fixture_slot", None)
        else:
            entries.append({"label_prefix": key, **approval})
    p = WORK / "teams.yaml"
    p.write_text(yaml.safe_dump({"teams": entries}, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return p


def _leaders(result) -> list[dict]:
    cfg = result.cfg
    groups: dict[str, dict] = {}
    counts: dict[str, Counter] = defaultdict(Counter)
    for rec in result.records:
        counts[rec.sales_leader][rec.status.value] += 1
    for label in sorted(counts):
        team = cfg.teams.lookup(label) if label else None
        key = team.label_prefix if team else label
        g = groups.setdefault(key, {"key": key, "labels": [], "status": team.status if team else "UNRESOLVED (not in teams.yaml)",
                                    "slot": team.slot if team and team.exportable else None,
                                    "describe": team.describe() if team else "No entry", "records": 0,
                                    "exported": 0})
        g["labels"].append(label or "(blank)")
        g["records"] += sum(counts[label].values())
        g["exported"] += counts[label].get(RecordStatus.EXPORTED.value, 0)
    return list(groups.values())


def _result(result) -> dict:
    cfg = result.cfg
    issues = Counter((i.code, i.severity.value) for r in result.records for i in r.issues)
    issues.update((i.code, i.severity.value) for i in result.response_issues)
    review = []
    for rec in sorted(result.records, key=lambda r: (r.source_row, r.block_index or 0)):
        flagged = [i for i in rec.issues if i.severity != Severity.INFO]
        if rec.status == RecordStatus.EXPORTED and not flagged:
            continue
        if len(review) < 300:
            review.append({"status": rec.status.value, "row": rec.source_row, "leader": rec.sales_leader,
                           "lead_source": rec.lead_source_original,
                           "reasons": ", ".join(dict.fromkeys(i.code for i in flagged))})
    owner = sum(1 for r in result.exported if r.slot_status == "OWNER_APPROVED_UNVERIFIED")
    return {
        "status": "ok", "batch_id": result.batch_id, "dry_run": result.dry_run, "counts": result.counts,
        "exported_verified": len(result.exported) - owner, "exported_owner_approved": owner,
        "checks": [{"name": n, "ok": ok} for n, ok, _ in result.checks],
        "mapping_status": [dict(zip(["leader", "status", "slot", "describe", "exported", "held"], row))
                           for row in mapping_status_rows(result)],
        "leaders": _leaders(result),
        "issues": [{"code": c, "severity": s, "count": n, "action": cfg.policy.action.get(c, "")}
                   for (c, s), n in sorted(issues.items(), key=lambda x: (_ORDER.index(x[0][1]), x[0][0]))],
        "review": review, "review_total": sum(1 for r in result.records if r.status != RecordStatus.EXPORTED
                                              or any(i.severity != Severity.INFO for i in r.issues)) + len(result.response_issues),
        "outputs": {k: str(v) for k, v in result.outputs.items()},
        "previous_batches": [{"batch_id": b.batch_id, "exported": b.records_exported} for b in result.previous_batches],
        "sheet": result.sheet.sheet_name, "ignored_sheets": result.sheet.ignored_sheets,
    }


def run(params_json: str) -> str:
    p = json.loads(params_json)
    try:
        out = WORK / "out"
        out.mkdir(parents=True, exist_ok=True)
        for old in out.glob("*"):
            old.unlink()
        opts = RunOptions(
            source=Path(p["source"]), template=Path(p["template"]), out_dir=out, config_dir=APP / "config",
            ledger_path=LEDGER, profiles_dir=WORK / "profiles",
            sellers_path=Path(p["sellers"]) if p.get("sellers") else None,
            teams_path=_teams_with_overrides(p.get("overrides") or {}),
            sheet=p.get("sheet") or None, last_id=p.get("last_id"), team_filter=p.get("team_filter") or None,
            leads_log=bool(p.get("leads_log")), dry_run=p.get("mode") != "export",
            timestamp_formats=[p["timestamp_format"]] if p.get("timestamp_format") else None,
            accept_suggestions=bool(p.get("accept_suggestions")), ignore_unmapped=bool(p.get("ignore_unmapped")))
        return json.dumps(_result(run_pipeline(opts)), default=str)
    except NeedsConfirmation as e:
        items = [{"column": i.letter, "header": i.header, "kind": i.kind, "target": i.target, "note": i.note}
                 for i in (e.cmap.needs_confirmation if e.cmap else [])]
        return json.dumps({"status": "needs_confirmation", "message": str(e), "sheet": e.sheet, "items": items})
    except _KNOWN as e:
        return json.dumps({"status": "error", "message": str(e)})
    except Exception as e:  # unexpected: show type and message, never workbook values
        return json.dumps({"status": "error", "message": f"Unexpected {type(e).__name__}: {e}",
                           "trace": traceback.format_exc(limit=3)})


def batches() -> str:
    return json.dumps([vars(b) for b in Ledger(LEDGER).list_batches()])


def void_batch(batch_id: str) -> str:
    try:
        return json.dumps({"status": "ok", "released": Ledger(LEDGER).void_batch(batch_id)})
    except KeyError as e:
        return json.dumps({"status": "error", "message": str(e)})


def classify(path: str) -> str:
    """Identify a dropped file by its header rows (not by file or sheet name). Returns JSON.

    kind: "source" (response export), "template" (MASTERLIST), "sellers" (seller directory CSV),
    "both" (a workbook holding both), or "unknown".
    """
    import csv
    import hashlib

    import openpyxl

    from fuel_parser.config import load_profile
    from fuel_parser.headers import match_key

    profile = load_profile(APP / "config" / "mapping_profile.yaml")
    src_keys = set(profile.source["signature_keys"])
    tpl_keys = {"completion time", "client's surname", "seller name1", "lead source"}
    digest = profile.dest["template_header_sha256"]
    p = Path(path)
    found = {"source_sheet": None, "template_sheet": None, "template_exact": False}

    def inspect(name: str, header: list) -> None:
        found["_sheets"] = found.get("_sheets", 0) + 1
        keys = {match_key(h) for h in header}
        if src_keys <= keys and found["source_sheet"] is None:
            found["source_sheet"] = name
        if tpl_keys <= keys:
            while header and header[-1] is None:
                header.pop()
            exact = hashlib.sha256("\x1f".join("" if v is None else str(v) for v in header).encode("utf-8")).hexdigest() == digest
            if found["template_sheet"] is None or (exact and not found["template_exact"]):
                found["template_sheet"], found["template_exact"] = name, exact

    try:
        if p.suffix.lower() == ".csv":
            with open(p, encoding="utf-8-sig", errors="replace", newline="") as fh:
                header = next(csv.reader(fh), [])
            if header and match_key(header[0]) == "seller_name":
                found.pop("_sheets", None)
                return json.dumps({"kind": "sellers", **found, "other_sheets": 0})
            inspect(p.stem, header)
        else:
            wb = openpyxl.load_workbook(p, read_only=True)
            try:
                for name in wb.sheetnames:
                    inspect(name, list(next(wb[name].iter_rows(min_row=1, max_row=1, values_only=True), ())))
            finally:
                wb.close()
    except Exception as e:  # unreadable file: report the type only, never content
        found.pop("_sheets", None)
        return json.dumps({"kind": "error", "message": type(e).__name__, **found, "other_sheets": 0})
    sheets_seen = found.pop("_sheets", 0)
    found["other_sheets"] = max(0, sheets_seen - 1) if found["source_sheet"] else 0
    kind = ("both" if found["source_sheet"] and found["template_sheet"] else "source" if found["source_sheet"]
            else "template" if found["template_sheet"] else "unknown")
    return json.dumps({"kind": kind, **found})
