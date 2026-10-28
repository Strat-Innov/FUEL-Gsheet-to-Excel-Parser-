"""Synthetic workbook builders. All names, e-mails and numbers are invented test data."""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

import openpyxl

TEMPLATE_HEADERS = [
    'Id', 'Start time', 'Completion time', 'Email', 'Name', 'SALES LEADER', 'Seller Name', "Seller's Email Address",
    'Seller Name1', "Seller's Email Address1", 'Seller Name2', "Seller's Email Address2", "Client's Surname",
    "Client's First Name", "Client's Middle Name\xa0", "Client's Mobile No.\n", "Client's Email Address",
    "Client's Residence Address", 'Project of Interest (First Choice)', 'Project of Interest (Second Choice)', 'REMARKS',
    'Lead Source', 'If\xa0Booth, please choose which\xa0Booth Location\xa0below:\n',
    'Name of\xa0existing buyer who referred\xa0the client to you\n\n', 'Question',
    'If Digital -\xa0Official Social Media, please choose below:\n', 'If Digital -\xa0Seller Social Media, please choose below:',
    'If\xa0Own Source, please choose below:\n', 'If\xa0Repeat Buyer, please choose from the listed projects below:\n',
    'If\xa0Roadshow, please specify below the\xa0name and date\xa0of roadshow event\xa0(e.g. Filinvest Alabang Roadshow - January 1, 2026)',
    'If\xa0Saturation, please specify\xa0exact location\xa0below\n', 'Question1',
    'If\xa0Sponsorship Event, please specify below the\xa0name and date\xa0of event\xa0(e.g. Cardiologist Conference 2026 - January 1, 2026)\n\n',
    "Source of Awareness\nCLIENT'S PERSPECTIVE:\xa0please choose the applicable source of awareness below\xa0(you may choose more than one)\n",
    '\nIf Digital -\xa0Seller Facebook Page,\xa0please choose below:',
]

BLOCK_FIELDS = ["seller_name", "surname", "first_name", "middle_name", "mobile", "client_email", "residence",
                "project_first", "project_second", "remarks"]


def block_headers(n: int) -> list[str]:
    """Google-style duplicate headers. Middle-name suffixes deliberately do not follow the block number."""
    s = "" if n == 1 else f" {n}"
    middle = {1: "Client's Middle Name", 2: "Client's Middle Name ", 3: "Client's Middle Name  2"}[n]
    return [f"Seller Name {s}" if n > 1 else "Seller Name ", f"Client's Surname{s}", f"Client's First Name{s}", middle,
            f"Mobile No. (Client){s}", f"Email Address (Client){s}", f"Residence Address (Client){s}",
            f"Project of Interest (First Choice){s}", f"Project of Interest (Second Choice){s}", f"REMARKS {s}" if n > 1 else "REMARKS "]


V1_DETAILS = [
    ("own_source_v1", "Own Source "), ("others_own_source_v1", "If others, please specify "),
    ("referral_v1", "Referral"), ("others_referral_v1", "If others, please specify  2"),
    ("digital_v1", "Digital"), ("others_digital_v1", "If others, please specify  3"),
    ("seller_fb_page_v1", "Seller Facebook Page "), ("saturation_v1", "If Saturation, please specify location"),
    ("booth_v1", "If Booth, please specify Booth Location"),
    ("sponsorship_v1", "If Sponsorship Event, please specify name of event & date (e.g. Cardiologist Conference 2026 - January 1, 2026) "),
    ("repeat_v1", "Projects owned/bought by Client "), ("others_repeat_project_v1", "If Others, please specify project"),
    ("others_lead_source_v1", "If others, please specify"),
]
V2_DETAILS = [
    ("booth", "If Booth, please choose which Booth Location below:"),
    ("referrer", "Name of existing buyer who referred the client to you"),
    ("corporate", "If Corporate, please choose below:"), ("others_corporate", "If Others, please specify below"),
    ("digital_official", "If Digital - Official Social Media, please choose below: "),
    ("digital_seller", "If Digital - Seller Social Media, please choose below: "),
    ("others_digital_seller", "If Others, please specify below 2"),
    ("seller_fb_page", "If Digital - Seller Facebook Page, please choose below:"),
    ("own_source", "If Own Source, please choose below:"), ("others_own_source", "If Others, please specify below 3"),
    ("repeat_buyer", "If Repeat Buyer, please choose from the listed projects below:"),
    ("others_repeat_buyer", "If Others, please specify below 4"),
    ("roadshow", "If Roadshow, please specify below the name and date of roadshow event (e.g. Filinvest Alabang Roadshow - January 1, 2026)"),
    ("saturation", "If Saturation, please specify exact location below"),
    ("showroom", "If Showroom Walk-In, please specify which Showroom below:"),
    ("sponsorship", "If Sponsorship Event, please specify below the name and date of event (e.g. Cardiologist Conference 2026 - January 1, 2026)"),
    ("others_lead_source", "If Others, please specify below 5"),
]


def default_layout() -> list[tuple[str, str]]:
    """(key, header) pairs in MARF-like order. Keys: response fields, b<n>.<field>, detail ids."""
    cols = [("timestamp", "Timestamp"), ("submitter", "Email Address"), ("leader", "SALES LEADER")]
    cols += [(f"b1.{f}", h) for f, h in zip(BLOCK_FIELDS, block_headers(1))]
    cols += [("lead_source", "Lead Source ")]
    cols += V1_DETAILS
    cols += [(f"b2.{f}", h) for f, h in zip(BLOCK_FIELDS, block_headers(2))]
    cols += V2_DETAILS
    cols += [("awareness", "Source of Awareness")]
    cols += [(f"b3.{f}", h) for f, h in zip(BLOCK_FIELDS, block_headers(3))]
    return cols


def response(block: int = 2, ts: object = datetime(2026, 6, 1, 10, 0, 0), submitter: str = "seller.one@example.com",
             leader: str = "TEAM ALPHA (AAD: Test Leader)", lead_source: str = "Own Source", **fields) -> dict:
    """One synthetic response. Block fields are given unprefixed and placed in `block`."""
    row = {"timestamp": ts, "submitter": submitter, "leader": leader, "lead_source": lead_source}
    client = {"seller_name": "AS: Test Seller", "surname": "Testsurname", "first_name": "Testfirst",
              "mobile": "09000000001", "residence": "METRO MANILA - Test City", "project_first": "Botanika"}
    for k, v in list(fields.items()):
        if k in BLOCK_FIELDS:
            client[k] = v
            fields.pop(k)
    if block:
        row.update({f"b{block}.{k}": v for k, v in client.items()})
    row.update(fields)
    return row


def write_source(path: Path, rows: list[dict], layout: list[tuple[str, str]] | None = None,
                 sheet: str = "MARF RAW", extra_sheets: dict[str, list[list]] | None = None) -> Path:
    layout = layout or default_layout()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append([h for _, h in layout])
    for i, r in enumerate(rows, start=2):
        for c, (k, _) in enumerate(layout, start=1):
            v = r.get(k)
            cell = ws.cell(i, c)
            cell.value = v
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"  # exports hold typed text, not formulas
    for name, data in (extra_sheets or {}).items():
        es = wb.create_sheet(name)
        for line in data:
            es.append(line)
    wb.save(path)
    return path


def write_source_csv(path: Path, rows: list[dict], layout: list[tuple[str, str]] | None = None,
                     ts_text=lambda d: d.strftime("%m/%d/%Y %H:%M:%S")) -> Path:
    layout = layout or default_layout()
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow([h for _, h in layout])
        for r in rows:
            out = []
            for k, _ in layout:
                v = r.get(k)
                out.append(ts_text(v) if isinstance(v, datetime) else ("" if v is None else v))
            w.writerow(out)
    return path


def write_template(path: Path, headers: list[str] | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MASTERLIST"
    ws.append(headers or TEMPLATE_HEADERS)
    wb.save(path)
    return path


TEAMS_TEST_YAML = """
teams:
  - {label_prefix: "TEAM ALPHA (AAD:", slot: 1, status: OWNER_APPROVED_UNVERIFIED, evidence_type: business_rule, approved_by: "Test Owner", approved_on: "2026-10-28"}
  - {label_prefix: "TEAM BETA (AAD:", slot: 2, status: VERIFIED, evidence_type: masterlist_export, evidence_ref: "test.xlsx sha256:abc"}
  - {label_prefix: "TEAM GAMMA (AAD:", slot: null, status: UNRESOLVED, evidence_type: none}
  - {label_prefix: "TEAM DELTA (AAD:", slot: 3, status: VERIFIED, evidence_type: owner_confirmation, evidence_ref: "e-mail 2026-10-01", destination_label: "TEAM DELTA (AAD: Renamed Leader)"}
"""
