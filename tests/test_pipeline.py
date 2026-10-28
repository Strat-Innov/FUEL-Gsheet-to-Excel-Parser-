"""End-to-end tests on synthetic workbooks (design G, M.6, N.6)."""

import hashlib
import socket
import sqlite3
from datetime import datetime

import openpyxl
import pytest

from fuel_parser.config import ConfigError
from fuel_parser.model import RecordStatus
from fuel_parser.pipeline import NeedsConfirmation
from fuel_parser.reader import ReadError
from fuel_parser.writer import TemplateError, leads_log_row, slot_mismatch

from .fixtures.builders import TEMPLATE_HEADERS, default_layout, response, write_source, write_source_csv


def recs(result):
    return sorted(result.records, key=lambda r: (r.source_row, r.block_index or 0))


def codes(rec):
    return set(rec.codes())


def import_rows(result):
    ws = openpyxl.load_workbook(result.outputs["import"])["MASTERLIST"]
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def balanced(result):
    return all(ok for _, ok, _ in result.checks)


def test_T1_normal_v2_response(env):
    r = env.run(env.source([response(own_source="Friend", awareness="Word of Mouth")]))
    (rec,) = recs(r)
    assert rec.status == RecordStatus.EXPORTED and balanced(r)
    v = rec.values
    assert (v["surname"], v["first_name"], v["mobile"], v["lead_source"], v["own_source"], v["awareness"]) == \
        ("Testsurname", "Testfirst", "09000000001", "Own Source", "Friend", "Word of Mouth")
    assert v["start_time"] == v["completion_time"] == datetime(2026, 6, 1, 10, 0)  # H-4
    assert v["email"] == "seller.one@example.com" and v["name"] is None  # H-3
    row = import_rows(r)[0]
    assert row[0] == 101 and len(row) == 35  # H-5: continues from --last-id


def test_T2_multi_block_response_is_held_not_merged(env):
    row = response(block=2)
    row.update({"b3.seller_name": "AS: Other Seller", "b3.surname": "Othersurname", "b3.first_name": "Other",
                "b3.mobile": "09000000002"})
    r = env.run(env.source([row]))
    assert len(r.records) == 2 and all(x.status == RecordStatus.REVIEW for x in r.records)
    assert all("MULTI_BLOCK_RESPONSE" in codes(x) for x in r.records)
    assert r.counts["responses_multi_block"] == 1 and balanced(r) and import_rows(r) == []


def test_T3_response_without_client_block(env):
    r = env.run(env.source([response(block=0), response()]))
    assert r.counts["responses_no_client"] == 1 and len(r.records) == 1 and balanced(r)
    assert r.response_issues[0].code == "NO_CLIENT_DATA"


def test_T4_slot_follows_sales_leader(env):
    r = env.run(env.source([response(leader="TEAM BETA (AAD: Someone)")]))
    v = recs(r)[0].values
    assert v["seller_name_2"] == "AS: Test Seller" and v["seller_email_2"] == "seller.one@example.com"
    assert not any(v[f"seller_{k}_{n}"] for k in ("name", "email") for n in (1, 3))


def test_T5_directory_email_differs_from_submitter(env):
    (env.config / "sellers.csv").write_text(
        "seller_name,proposed_email,status,reviewer_email\nAS: Test Seller,dir@example.com,APPROVED,\n", encoding="utf-8")
    rec = recs(env.run(env.source([response()])))[0]
    assert rec.values["seller_email_1"] == "dir@example.com" and "SUBMITTER_NOT_SELLER" in codes(rec)


def test_unapproved_directory_rows_fall_back_to_submitter(env):
    (env.config / "sellers.csv").write_text(
        "seller_name,proposed_email,status,reviewer_email\nAS: Test Seller,dir@example.com,DRAFT - NOT APPROVED,\n", encoding="utf-8")
    rec = recs(env.run(env.source([response()])))[0]
    assert rec.values["seller_email_1"] == "seller.one@example.com" and "SELLER_EMAIL_FROM_SUBMITTER" in codes(rec)


def test_T6_missing_seller_name_writes_no_email(env):
    rec = recs(env.run(env.source([response(seller_name=None)])))[0]
    assert "SELLER_NAME_MISSING" in codes(rec) and rec.status == RecordStatus.REVIEW
    assert not any(rec.values[f"seller_{k}_{n}"] for k in ("name", "email") for n in (1, 2, 3))


def test_T7_missing_seller_email(env):
    rec = recs(env.run(env.source([response(submitter=None)])))[0]
    assert rec.values["seller_name_1"] == "AS: Test Seller" and rec.values["seller_email_1"] is None
    assert "SELLER_EMAIL_MISSING" in codes(rec) and rec.status == RecordStatus.EXPORTED


def test_T8_T46_unknown_and_unresolved_teams_are_held(env):
    r = env.run(env.source([response(leader="TEAM NEW (AAD: X)"), response(leader="TEAM GAMMA (AAD: Y)")]))
    for rec in r.records:
        assert rec.status == RecordStatus.REVIEW and "SLOT_UNRESOLVED" in codes(rec)
        assert not any(rec.values[f"seller_{k}_{n}"] for k in ("name", "email") for n in (1, 2, 3))
    assert import_rows(r) == [] and balanced(r)


def test_T9_T23_moved_columns_give_identical_records(env):
    rows = [response(own_source="Friend"), response(block=3, ts=datetime(2026, 6, 2, 9, 0), surname="Second")]
    base = env.run(env.source(rows, name="a.xlsx"), dry_run=True)
    layout = default_layout()
    resp_cols, rest = layout[:3], layout[3:]
    b1, after = rest[:10], rest[10:]
    moved = after + b1 + resp_cols[::-1]  # blocks reordered, response fields moved to the end
    other = env.run(env.source(rows, name="b.xlsx", layout=moved), dry_run=True)
    key = lambda res: sorted((r.values["surname"], tuple(sorted((k, str(v)) for k, v in r.values.items()))) for r in res.records)
    assert key(base) == key(other)


def test_T12_blank_optional_fields_stay_blank(env):
    rec = recs(env.run(env.source([response()])))[0]
    for f in ("middle_name", "client_email", "project_second", "remarks", "awareness", "question"):
        assert rec.values[f] is None


def test_T13_null_tokens_blanked_and_logged(env):
    rec = recs(env.run(env.source([response(client_email="N/A", middle_name="None")])))[0]
    assert rec.values["client_email"] is None and rec.values["middle_name"] is None
    assert "NULL_TOKEN_BLANKED" in codes(rec)


def test_T15_T16_T31_mobile_written_as_text(env):
    rows = [response(mobile="09000001234"), response(ts=datetime(2026, 6, 1, 11), mobile=9000001234, surname="B"),
            response(ts=datetime(2026, 6, 1, 12), mobile="9999999999", surname="C")]
    r = env.run(env.source(rows))
    ws = openpyxl.load_workbook(r.outputs["import"])["MASTERLIST"]
    cells = [ws.cell(i, 16) for i in (2, 3, 4)]
    assert [c.value for c in cells] == ["09000001234", "9000001234", "9999999999"]
    assert all(c.data_type == "s" and c.number_format == "@" for c in cells)
    a, b, c = recs(r)
    assert "MOBILE_FORMAT_UNEXPECTED" not in codes(a)
    assert {"MOBILE_WAS_NUMERIC", "MOBILE_FORMAT_UNEXPECTED"} <= codes(b) and b.status == RecordStatus.EXPORTED
    assert "repeated_digits" in next(i.message for i in c.issues if i.code == "MOBILE_FORMAT_UNEXPECTED")
    assert c.status == RecordStatus.EXPORTED  # warnings do not block (D.8)


def test_T17_same_file_twice_exports_nothing_new(env):
    src = env.source([response(), response(ts=datetime(2026, 6, 2), surname="Other")])
    first = env.run(src)
    second = env.run(src)
    assert first.counts["records_exported"] == 2
    assert second.counts["records_exported"] == 0 and second.counts["records_already_imported"] == 2
    assert second.previous_batches and balanced(second)


def test_T18_newer_export_only_new_rows(env):
    old = [response(ts=datetime(2026, 6, d), surname=f"S{d}") for d in range(1, 4)]
    env.run(env.source(old, name="v1.xlsx"))
    new = old + [response(ts=datetime(2026, 7, d), surname=f"N{d}") for d in range(1, 6)]
    r = env.run(env.source(new, name="v2.xlsx"))
    assert r.counts["records_exported"] == 5 and r.counts["records_already_imported"] == 3


def test_T19_exact_content_duplicate_is_held(env):
    rows = [response(ts=datetime(2026, 6, 1)), response(ts=datetime(2026, 6, 20))]
    r = env.run(env.source(rows))
    first, second = recs(r)
    assert first.status == RecordStatus.EXPORTED
    assert second.status == RecordStatus.DUPLICATE_HELD and "PROBABLE_DUPLICATE" in codes(second)


def test_T20_same_client_different_seller_both_exported(env):
    rows = [response(), response(ts=datetime(2026, 6, 5), seller_name="AS: Another Seller", project_first="The Levels")]
    a, b = recs(env.run(env.source(rows)))
    assert a.status == b.status == RecordStatus.EXPORTED
    assert "POSSIBLE_DUPLICATE_CLIENT" in codes(a) and a.duplicate_group == b.duplicate_group != ""


def test_T21_renamed_header_needs_confirmation(env):
    layout = [(k, "Mobile Number (Client) 2" if k == "b2.mobile" else h) for k, h in default_layout()]
    layout = [(k, "If Own Source, please pick below:" if k == "own_source" else h) for k, h in layout]
    src = env.source([response(own_source="Friend")], layout=layout)
    with pytest.raises(Exception):  # block structure is checked strictly: a renamed block field stops the run
        env.run(src)
    layout2 = [(k, "If Own Source, please pick below:" if k == "own_source" else h) for k, h in default_layout()]
    src2 = env.source([response(own_source="Friend")], name="s2.xlsx", layout=layout2)
    with pytest.raises(NeedsConfirmation) as e:
        env.run(src2)
    assert e.value.cmap.needs_confirmation[0].kind == "SUGGESTED"
    r = env.run(src2, accept_suggestions=True)
    assert recs(r)[0].values["own_source"] == "Friend"
    assert env.run(src2, dry_run=True).counts["records_built"] == 1  # decision saved for this header signature


def test_T22_unrecognised_non_empty_column_must_be_confirmed(env):
    layout = default_layout() + [("budget", "Budget Range")]
    src = env.source([response(budget="5M")], layout=layout)
    with pytest.raises(NeedsConfirmation):
        env.run(src)
    r = env.run(src, ignore_unmapped=True)
    assert r.counts["records_built"] == 1


def test_T24_v1_lead_source_translation(env):
    rows = [response(block=1, lead_source="Ownsource", own_source_v1="Friend"),
            response(block=1, ts=datetime(2026, 6, 2), surname="D", lead_source="Digital", digital_v1="Facebook Marketplace"),
            response(block=1, ts=datetime(2026, 6, 3), surname="E", lead_source="Digital", digital_v1="Project's Main Page"),
            response(block=1, ts=datetime(2026, 6, 4), surname="F", lead_source="Booth", booth_v1="Westgate")]
    a, b, c, d = recs(env.run(env.source(rows)))
    assert (a.values["lead_source"], a.values["own_source"]) == ("Own Source", "Friend")
    assert (b.values["lead_source"], b.values["digital_seller"]) == ("Digital - Seller Social Media", "Facebook Marketplace")
    assert (c.values["lead_source"], c.values["digital_official"]) == ("Digital - Official Social Media", "Project's Main Page")
    assert (d.values["lead_source"], d.values["booth"]) == ("Booth Manning", "Westgate")
    assert a.lead_source_original == "Ownsource" and all(x.status == RecordStatus.EXPORTED for x in (a, b, c, d))


def test_T25_unapproved_legacy_values_are_held(env):
    rows = [response(block=1, lead_source="Site Walk-in"),
            response(block=1, ts=datetime(2026, 6, 2), surname="R", lead_source="Referral", referral_v1="Friend"),
            response(block=1, ts=datetime(2026, 6, 3), surname="O", lead_source="Digital", digital_v1="Others")]
    for rec in env.run(env.source(rows)).records:
        assert rec.status == RecordStatus.REVIEW and "UNMAPPED_VALUE" in codes(rec)


def test_H8_others_text_appended_to_remarks_with_label(env):
    rec = recs(env.run(env.source([response(remarks="Call after 5", own_source="Others", others_own_source="Old contact")])))[0]
    assert rec.values["remarks"] == "Call after 5 | [Own Source – Others] Old contact"
    assert "ORPHAN_OTHERS_TEXT" not in codes(rec)


def test_T26_orphan_others_text_is_kept_and_flagged(env):
    rec = recs(env.run(env.source([response(own_source="Friend", others_own_source="Extra text")])))[0]
    assert rec.values["remarks"] == "[Own Source – Others] Extra text" and "ORPHAN_OTHERS_TEXT" in codes(rec)


def test_lead_source_others_text(env):
    rec = recs(env.run(env.source([response(lead_source="Others", others_lead_source="Walk-by")])))[0]
    assert rec.values["remarks"] == "[Lead Source – Others] Walk-by"


def test_T27_formula_like_text_is_literal(env):
    r = env.run(env.source([response(remarks="=SUM(A1)")]))
    cell = openpyxl.load_workbook(r.outputs["import"])["MASTERLIST"].cell(2, 21)
    assert cell.value == "=SUM(A1)" and cell.data_type == "s"


def test_T28_header_bytes_equal_template(env):
    r = env.run(env.source([response()]))
    ws = openpyxl.load_workbook(r.outputs["import"])["MASTERLIST"]
    assert [c.value for c in ws[1]] == TEMPLATE_HEADERS


def test_T29_inputs_unchanged(env):
    src = env.source([response()])
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (src, env.template)]
    env.run(src, leads_log=True)
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (src, env.template)]


def test_T30_void_batch_releases_records(env):
    from fuel_parser.history import Ledger
    src = env.source([response()])
    first = env.run(src)
    assert env.run(src).counts["records_exported"] == 0
    assert Ledger(env.ledger).void_batch(first.batch_id) == 1
    assert env.run(src).counts["records_exported"] == 1


def test_T32_invalid_seller_email_kept_and_warned(env):
    rec = recs(env.run(env.source([response(submitter="name@gmail,com")])))[0]
    assert rec.values["seller_email_1"] == "name@gmail,com" and "INVALID_EMAIL" in codes(rec)
    assert rec.status == RecordStatus.EXPORTED


def test_T33_misspelled_domain_not_flagged(env):
    rec = recs(env.run(env.source([response(submitter="name@gmaol.com", client_email="x@company.ph")])))[0]
    assert "INVALID_EMAIL" not in codes(rec)


def test_T34_slot_mismatch_detected_without_moving():
    values = {"seller_name_2": "AS: Seller", "seller_email_2": "s@example.com"}
    assert slot_mismatch(values, 1) is True and slot_mismatch(values, 2) is False
    assert values == {"seller_name_2": "AS: Seller", "seller_email_2": "s@example.com"}


def test_T35_leads_log_projection(env):
    rows = [response(lead_source="Digital - Seller Social Media", digital_seller="Seller Facebook Page",
                     seller_fb_page="Facebook Page with Company's Boosting Budget", awareness="Events", mobile="9999999999"),
            response(ts=datetime(2026, 6, 2), surname="W", lead_source="Showroom Walk-In", showroom="Test Showroom"),
            response(ts=datetime(2026, 6, 3), surname="O", lead_source="Digital - Official Website")]
    r = env.run(env.source(rows), leads_log=True)
    ws = openpyxl.load_workbook(r.outputs["leads_log"])["LEADS LOG"]
    assert [c.value for c in ws[1]][0] == "TIMESTAMP " and ws.max_column == 17
    a, b, c = (list(row) for row in ws.iter_rows(min_row=2, values_only=True))
    assert a[3] == "AS: Test Seller" and a[2] == "TEAM ALPHA (AAD: Test Leader)"  # AS = seller, AAD = leader
    assert a[7] == "9999999999" and a[14] == "Seller Facebook Page" and a[15] is None and a[16] == "Events"
    assert b[14] == "Test Showroom" and c[14] is None
    assert ws.cell(2, 1).number_format == "mmm/dd/yyyy hh:mm:ss" and isinstance(ws.cell(2, 1).value, datetime)


def test_leads_log_row_uses_configured_slot_only(env):
    from fuel_parser.config import load_profile
    prof = load_profile(env.config / "mapping_profile.yaml")
    values = {"seller_name_1": "AS: One", "seller_email_1": "one@example.com", "lead_source": "Own Source", "own_source": "Friend"}
    assert leads_log_row(values, 2, prof)[1] is None and leads_log_row(values, 1, prof)[3] == "AS: One"


def test_T36_csv_matches_xlsx(env):
    rows = [response(own_source="Friend", awareness="Word of Mouth, Events"), response(ts=datetime(2026, 6, 14, 9, 30), surname="Z")]
    x = env.run(env.source(rows), dry_run=True)
    c = env.run(write_source_csv(env.root / "s.csv", rows), dry_run=True, timestamp_formats=["%m/%d/%Y %H:%M:%S"])
    strip = lambda res: [{k: v for k, v in r.values.items()} for r in recs(res)]
    assert strip(x) == strip(c)
    assert recs(c)[0].values["mobile"] == "09000000001"


def test_T37_csv_ambiguous_dates(env):
    rows = [response(ts=datetime(2026, 4, 3, 10, 0))]
    src = write_source_csv(env.root / "amb.csv", rows)
    rec = recs(env.run(src, dry_run=True))[0]
    assert "AMBIGUOUS_DATE" in codes(rec) and rec.status == RecordStatus.REVIEW
    ok = recs(env.run(src, dry_run=True, timestamp_formats=["%m/%d/%Y %H:%M:%S"]))[0]
    assert ok.timestamp == datetime(2026, 4, 3, 10, 0)


def test_T38_single_response_sheet_used_automatically_others_ignored(env):
    team_sheet = [["2026 SALES FUNNEL REPORT"], [], [], [None, "TIMESTAMP ", "SELLER EMAIL ADDRESS", "AAD", "AS"]]
    src = env.source([response()], sheet="Form Responses 1", extra_sheets={"TEAM ALPHA": team_sheet, "Notes": [["x"]]})
    r = env.run(src)
    assert r.sheet.sheet_name == "Form Responses 1" and r.sheet.ignored_sheets == ["TEAM ALPHA", "Notes"]
    summary = dict(row[:2] for row in openpyxl.load_workbook(r.outputs["audit"])["Summary"].iter_rows(values_only=True))
    assert summary["Response sheet used"] == "Form Responses 1" and summary["Other sheets ignored"] == "TEAM ALPHA, Notes"


def test_marf_raw_preferred_over_other_response_like_sheets(env):
    other = [[h for _, h in default_layout()]]
    src = env.source([response()], extra_sheets={"Copy of responses": other})
    assert env.run(src).sheet.sheet_name == "MARF RAW"


def test_several_response_sheets_without_marf_raw_need_a_choice(env):
    other = [[h for _, h in default_layout()]]
    src = env.source([response()], sheet="Responses A", extra_sheets={"Responses B": other})
    with pytest.raises(NeedsConfirmation) as e:
        env.run(src)
    assert e.value.sheet == "Responses A"
    assert env.run(src, sheet="Responses A").counts["records_built"] == 1


def test_team_sheet_check_is_opt_in(env):
    src = env.source([response()], extra_sheets={"TEAM ALPHA": [[None, "TIMESTAMP", "AAD", "AS"]]})
    assert "TeamSheetCheck" not in openpyxl.load_workbook(env.run(src, dry_run=True).outputs["audit"]).sheetnames
    assert "TeamSheetCheck" in openpyxl.load_workbook(env.run(src, dry_run=True, team_sheet_check=True).outputs["audit"]).sheetnames


def test_T39_team_sheets_only_is_rejected(env):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEAM ALPHA"
    ws.append(["2026 SALES FUNNEL REPORT"])
    ws.append([None, "TIMESTAMP ", "SELLER EMAIL ADDRESS", "AAD", "AS"])
    p = env.root / "teams_only.xlsx"
    wb.save(p)
    with pytest.raises(ReadError, match="No response sheet found"):
        env.run(p)


def test_T43_altered_template_stops_before_output(env):
    from .fixtures.builders import write_template
    bad = write_template(env.root / "bad.xlsx", TEMPLATE_HEADERS[:-1] + ["Changed"])
    with pytest.raises(TemplateError):
        env.run(env.source([response()]), template=bad)
    assert not env.out.exists() or not any(env.out.glob("FUEL_*"))


def test_T44_runs_with_networking_disabled(env, monkeypatch):
    def blocked(*a, **k):
        raise AssertionError("network access attempted")
    monkeypatch.setattr(socket, "socket", blocked)
    monkeypatch.setattr(socket, "create_connection", blocked)
    assert env.run(env.source([response()])).counts["records_exported"] == 1


def test_T45_T48_owner_approved_is_never_reported_as_verified(env):
    r = env.run(env.source([response(), response(ts=datetime(2026, 6, 2), surname="B", leader="TEAM BETA (AAD: Z)")]))
    a, b = recs(r)
    assert "SLOT_OWNER_APPROVED_UNVERIFIED" in codes(a) and a.status == RecordStatus.EXPORTED
    assert "SLOT_OWNER_APPROVED_UNVERIFIED" not in codes(b)
    wb = openpyxl.load_workbook(r.outputs["audit"])
    status = {row[0]: row for row in wb["MappingStatus"].iter_rows(min_row=2, values_only=True)}
    alpha = status["TEAM ALPHA (AAD: Test Leader)"]
    assert alpha[1] == "OWNER_APPROVED_UNVERIFIED" and "not verified" in alpha[3] and not alpha[3].startswith("Verified")
    assert status["TEAM BETA (AAD: Z)"][3].startswith("Verified")
    headline = next(row[1] for row in wb["Summary"].iter_rows(values_only=True) if row[0] == "Headline")
    assert "Exported: 1 verified · 1 owner-approved (unverified)" in headline
    review = [row for row in wb["Review"].iter_rows(min_row=2, values_only=True)]
    assert any("SLOT_OWNER_APPROVED_UNVERIFIED" in row[1] for row in review)


def test_T50_destination_label(env):
    a, b = recs(env.run(env.source([response(leader="TEAM DELTA (AAD: Original)"),
                                    response(ts=datetime(2026, 6, 2), surname="B", leader="TEAM BETA (AAD: Kept)")])))
    assert a.values["sales_leader"] == "TEAM DELTA (AAD: Renamed Leader)" and a.values["seller_name_3"]
    assert b.values["sales_leader"] == "TEAM BETA (AAD: Kept)"


def test_T51_corporate_value_is_held_and_question_blank(env):
    rec = recs(env.run(env.source([response(lead_source="Corporate", corporate="Some Corp")])))[0]
    assert rec.values["question"] is None and "UNMAPPED_COLUMN" in codes(rec) and rec.status == RecordStatus.REVIEW


def test_orphan_detail_warning(env):
    rec = recs(env.run(env.source([response(lead_source="Saturation", saturation="Test Mall", booth="Westgate")])))[0]
    assert "ORPHAN_DETAIL" in codes(rec) and rec.values["booth"] == "Westgate"


def test_projects_multi_first_choice_and_unknown(env):
    rec = recs(env.run(env.source([response(project_first="Studio N, The Levels", project_second="Gold Ridge")])))[0]
    assert rec.values["project_first"] == "Studio N;The Levels" and rec.values["project_second"] == "Gold Ridge"
    assert {"MULTI_FIRST_CHOICE", "UNKNOWN_PROJECT"} <= codes(rec)


def test_numeric_remark_keeps_digits(env):
    rec = recs(env.run(env.source([response(remarks=9000001234)])))[0]
    assert rec.values["remarks"] == "9000001234"


def test_team_filter_excludes_and_balances(env):
    r = env.run(env.source([response(), response(ts=datetime(2026, 6, 2), surname="B", leader="TEAM BETA (AAD: Z)")]),
                team_filter=["TEAM BETA"])
    assert r.counts["records_excluded_by_filter"] == 1 and r.counts["records_exported"] == 1 and balanced(r)


def test_last_id_required_when_exporting(env):
    with pytest.raises(ValueError, match="last Id"):
        env.run(env.source([response()]), last_id=None)


def test_dry_run_writes_no_import_and_no_ledger(env):
    src = env.source([response()])
    r = env.run(src, dry_run=True)
    assert "import" not in r.outputs and env.run(src).counts["records_exported"] == 1


def test_ledger_stores_no_client_data(env):
    env.run(env.source([response(surname="Uniquesurname", mobile="09000009999")]))
    dump = "\n".join(sqlite3.connect(env.ledger).iterdump())
    assert "Uniquesurname" not in dump and "09000009999" not in dump and "seller.one@example.com" not in dump


def test_audit_accounts_for_every_row(env):
    rows = [response(), response(block=0, ts=datetime(2026, 6, 2)), response(ts=datetime(2026, 6, 3), surname="X", leader="TEAM GAMMA (AAD: Q)")]
    r = env.run(env.source(rows))
    rowlog = list(openpyxl.load_workbook(r.outputs["audit"])["RowLog"].iter_rows(min_row=2, values_only=True))
    assert sorted(x[0] for x in rowlog) == [2, 3, 4] and balanced(r)
    statuses = {x[0]: x[2] for x in rowlog}
    assert statuses == {2: "EXPORTED", 3: "REVIEW", 4: "REVIEW"}


def test_config_error_for_unknown_policy_code(env):
    p = env.config / "validation_policy.yaml"
    p.write_text(p.read_text(encoding="utf-8").replace("  SLOT_UNRESOLVED:", "  SLOT_UNRESOLVED_X:"), encoding="utf-8")
    with pytest.raises(ConfigError, match="SLOT_UNRESOLVED"):
        env.run(env.source([response(leader="TEAM GAMMA (AAD: Q)")]))


def test_template_found_by_header_when_sheet_is_renamed(env):
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    wb.active.append(["not the template"])
    ws = wb.create_sheet("Sheet1")
    ws.append(TEMPLATE_HEADERS)
    p = env.root / "renamed_template.xlsx"
    wb.save(p)
    r = env.run(env.source([response()]), template=p)
    out = openpyxl.load_workbook(r.outputs["import"])
    assert out.sheetnames == ["MASTERLIST"] and [c.value for c in out["MASTERLIST"][1]] == TEMPLATE_HEADERS
